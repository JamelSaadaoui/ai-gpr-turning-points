"""Create a comfortable reading edition of the article without changing content."""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "paper" / "Geopolitical_Turning_Points_and_Oil_Prices_V3.docx"
TARGET = ROOT / "paper" / "Geopolitical_Turning_Points_and_Oil_Prices_V3_Readable.docx"
FONT = "Palatino Linotype"


def apply_font(run, size):
    run.font.name = FONT
    run.font.size = Pt(size)
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    for key in ("ascii", "hAnsi", "cs", "eastAsia"):
        fonts.set(qn("w:" + key), FONT)


doc = Document(SOURCE)
section = doc.sections[0]
section.left_margin = Inches(0.9)
section.right_margin = Inches(0.9)
section.top_margin = Inches(0.85)
section.bottom_margin = Inches(0.85)

normal = doc.styles["Normal"]
normal.font.name = FONT
normal.font.size = Pt(12)
normal.paragraph_format.line_spacing = 1.25
normal.paragraph_format.space_after = Pt(8)

for name, size in (("Title", 20), ("Heading 1", 15), ("Heading 2", 13), ("Caption", 11)):
    style = doc.styles[name]
    style.font.name = FONT
    style.font.size = Pt(size)
    style.paragraph_format.space_before = Pt(10)
    style.paragraph_format.space_after = Pt(6)

# Body prose and notes. Table text is handled separately below.
for paragraph in doc.paragraphs:
    in_table = bool(paragraph._p.xpath("ancestor::w:tbl"))
    if in_table:
        continue
    text = paragraph.text.strip()
    if paragraph.style.name == "Title":
        size = 20
    elif paragraph.style.name == "Heading 1":
        size = 15
    elif paragraph.style.name == "Heading 2":
        size = 13
    elif paragraph.style.name == "Caption":
        size = 11
    elif text.startswith("Notes:"):
        size = 10.5
        paragraph.paragraph_format.line_spacing = 1.15
        paragraph.paragraph_format.space_after = Pt(11)
    elif text.startswith("References"):
        size = 12
    else:
        size = 12
        paragraph.paragraph_format.line_spacing = 1.25
        paragraph.paragraph_format.space_after = Pt(8)
    for run in paragraph.runs:
        apply_font(run, size)
    if text.startswith("Table 6  Conditional strength"):
        paragraph.paragraph_format.page_break_before = True

# Tables remain compact enough to fit the page, but no text is below 10 pt.
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.line_spacing = 1.08
                paragraph.paragraph_format.space_before = Pt(4)
                paragraph.paragraph_format.space_after = Pt(4)
                for run in paragraph.runs:
                    apply_font(run, 10)

# Enlarge the page number slightly and retain centered placement.
for section in doc.sections:
    for paragraph in section.footer.paragraphs:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in paragraph.runs:
            apply_font(run, 10)

doc.core_properties.author = "Jamel Saadaoui"
doc.core_properties.last_modified_by = "Jamel Saadaoui"
doc.core_properties.comments = "Readable manuscript edition"
doc.save(TARGET)
print(TARGET)
