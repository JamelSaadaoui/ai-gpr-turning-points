"""Read-only checks of the spreadsheet companion and notebook structure."""

from pathlib import Path
import ast
import json
import re
import unittest
import zipfile
import xml.etree.ElementTree as ET

import numpy as np
import pandas as pd
import openpyxl
from PIL import Image

PACKAGE = Path(__file__).resolve().parents[1]
XLSX = PACKAGE / "data" / "Geopolitical_Turning_Points_Data.xlsx"


class SpreadsheetTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.raw = pd.read_csv(PACKAGE / "data" / "monthly_data.csv", parse_dates=["date"])
        cls.cached = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)
        cls.formulas = openpyxl.load_workbook(XLSX, read_only=True, data_only=False)

    @classmethod
    def tearDownClass(cls):
        cls.cached.close()
        cls.formulas.close()

    def test_workbook_raw_data_equals_csv(self):
        rows = list(self.cached["Monthly data"].values)
        actual = pd.DataFrame(rows[1:], columns=rows[0])
        self.assertEqual(list(actual.columns), list(self.raw.columns))
        pd.testing.assert_series_equal(pd.to_datetime(actual.date), self.raw.date)
        np.testing.assert_allclose(actual.iloc[:, 1:].to_numpy(float),
                                   self.raw.iloc[:, 1:].to_numpy(float),
                                   rtol=1e-13, atol=1e-13, equal_nan=True)

    def test_cached_spreadsheet_transforms_match_independent_calculation(self):
        rows = list(self.cached["Transformations"].values)
        actual = pd.DataFrame(rows[1:], columns=rows[0])
        expected = pd.DataFrame({"date": self.raw.date})
        events = list(self.raw.columns[5:])
        for event in events:
            expected[f"log1p_{event}"] = np.log1p(self.raw[event])
        for event in events:
            expected[f"d2_{event}"] = self.raw[event] - 2 * self.raw[event].shift(1) + self.raw[event].shift(2)
        expected["d_lgop"] = self.raw.lgop.diff()
        self.assertEqual(list(actual.columns), list(expected.columns))
        pd.testing.assert_series_equal(pd.to_datetime(actual.date), expected.date)
        numeric = actual.iloc[:, 1:].replace("", np.nan).apply(pd.to_numeric)
        np.testing.assert_allclose(numeric.to_numpy(float), expected.iloc[:, 1:].to_numpy(float),
                                   rtol=1e-12, atol=1e-12, equal_nan=True)

    def test_transformation_cells_are_formulas_and_error_free(self):
        for row in self.formulas["Transformations"].iter_rows(min_row=2):
            for cell in row:
                self.assertEqual(cell.data_type, "f", f"Expected formula at {cell.coordinate}")
        for sheet in self.cached:
            for row in sheet:
                for cell in row:
                    self.assertNotEqual(cell.data_type, "e", f"Excel error in {sheet.title}/{cell.coordinate}")

    def test_dictionary_freezes_and_no_external_workbook_links(self):
        self.assertEqual(self.cached.sheetnames, ["Monthly data", "Transformations", "Dictionary"])
        dictionary_text = " ".join(str(value) for row in self.cached["Dictionary"].values for value in row if value)
        for expected in ("No splice", "October 2025", "https://", "lag"):
            self.assertIn(expected, dictionary_text)
        ns = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
        with zipfile.ZipFile(XLSX) as archive:
            self.assertFalse(any(name.startswith("xl/externalLinks/") for name in archive.namelist()))
            for name in ("xl/worksheets/sheet1.xml", "xl/worksheets/sheet2.xml"):
                tree = ET.fromstring(archive.read(name))
                pane = tree.find(".//s:pane", ns)
                self.assertIsNotNone(pane)
                self.assertEqual(pane.attrib["state"], "frozen")


class NotebookStructureTests(unittest.TestCase):
    def test_notebook_schema_and_python_syntax(self):
        notebook = json.loads((PACKAGE / "Geopolitical_Turning_Points_Replication.ipynb").read_text())
        self.assertEqual(notebook["nbformat"], 4)
        self.assertEqual(notebook["metadata"]["kernelspec"]["language"], "python")
        ids = [cell["id"] for cell in notebook["cells"]]
        self.assertEqual(len(ids), len(set(ids)))
        self.assertTrue(all(re.fullmatch(r"[a-zA-Z0-9_-]{1,64}", item) for item in ids))
        for i, cell in enumerate(notebook["cells"]):
            source = "".join(cell["source"])
            self.assertNotIn("/workspace/", source)
            if cell["cell_type"] == "code":
                ast.parse(source, filename=f"notebook_cell_{i}")
                self.assertIn("outputs", cell)
                self.assertIn("execution_count", cell)

    def test_publication_images_are_complete(self):
        figures = PACKAGE / "figures"
        expected = ("fig01_all_variables", "fig02_turning_points",
                    "fig03_brent_separate_irfs", "fig03_wti_separate_irfs",
                    "fig04_joint_irfs", "figS01_mechanism_irfs")
        for name in expected:
            with self.subTest(figure=name):
                with Image.open(figures / f"{name}.png") as picture:
                    self.assertGreater(picture.width, 2000)
                    self.assertGreater(picture.height, 2000)
                    picture.verify()
                self.assertTrue((figures / f"{name}.pdf").read_bytes().startswith(b"%PDF"))
                ET.parse(figures / f"{name}.svg")


if __name__ == "__main__":
    unittest.main()
